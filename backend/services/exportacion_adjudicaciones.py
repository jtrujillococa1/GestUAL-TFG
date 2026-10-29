"""Generación segura de listados exportables de la aplicación.

Las rutas resuelven antes el ámbito visible y este módulo solo transforma esas
filas autorizadas en Excel o PDF. Así se evita tanto exportar borradores como
filtrar datos en el navegador.
"""
from __future__ import annotations

from datetime import datetime
from html import escape
from io import BytesIO
from typing import Any, Iterable, Sequence


ENCABEZADOS_ADJUDICACIONES = (
    "Departamento", "Área", "Asignatura", "Código", "Grupo", "Tipo",
    "Cuatrimestre", "Horario", "Profesor/a", "Horas confirmadas",
)


def _texto(valor: Any) -> str:
    return "" if valor is None else str(valor)


def _filas_adjudicaciones(filas: Iterable[dict[str, Any]]) -> list[list[Any]]:
    return [
        [
            _texto(fila.get("departamento")), _texto(fila.get("area")),
            _texto(fila.get("asignatura")), _texto(fila.get("codigo")),
            _texto(fila.get("grupo")), _texto(fila.get("tipo")),
            _texto(fila.get("cuatrimestre")), _texto(fila.get("horario")),
            _texto(fila.get("profesor")), float(fila.get("horas_confirmadas", 0)),
        ]
        for fila in filas
    ]


def generar_excel_tabular(
    titulo: str,
    nombre_hoja: str,
    encabezados: Sequence[str],
    filas: Iterable[Sequence[Any]],
    anchos: Sequence[float],
    resumen: Sequence[tuple[str, Any]] = (),
    formatos_numericos: dict[int, str] | None = None,
) -> bytes:
    """Crea un XLSX simple y compatible con Excel de escritorio.

    Usa el autofiltro propio de la hoja, sin tablas OOXML. Así se elimina el
    XML de tabla que Excel tuvo que reparar en la descarga anterior.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    datos = [list(fila) for fila in filas]
    libro = Workbook()
    hoja_resumen = libro.active
    hoja_resumen.title = "Resumen"
    hoja_resumen.sheet_view.showGridLines = False
    hoja_resumen["A1"] = titulo
    hoja_resumen["A1"].font = Font(name="Arial", bold=True, size=15, color="173F6B")
    hoja_resumen["A3"] = "Generado"
    hoja_resumen["B3"] = datetime.now().strftime("%d/%m/%Y %H:%M")
    for indice, (etiqueta, valor) in enumerate(resumen, start=4):
        hoja_resumen.cell(indice, 1, etiqueta)
        hoja_resumen.cell(indice, 2, valor)
    for fila in range(3, 4 + len(resumen)):
        hoja_resumen.cell(fila, 1).font = Font(name="Arial", bold=True)
    hoja_resumen.column_dimensions["A"].width = 28
    hoja_resumen.column_dimensions["B"].width = 28

    hoja = libro.create_sheet(nombre_hoja)
    hoja.sheet_view.showGridLines = False
    hoja.append(list(encabezados))
    for fila in datos:
        hoja.append(fila)
    azul = PatternFill("solid", fgColor="1F4E78")
    borde = Border(bottom=Side(style="thin", color="D6E0EA"))
    for celda in hoja[1]:
        celda.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
        celda.fill = azul
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    hoja.row_dimensions[1].height = 28
    for fila in hoja.iter_rows(min_row=2):
        for celda in fila:
            celda.font = Font(name="Arial", size=10)
            celda.alignment = Alignment(vertical="top", wrap_text=True)
            celda.border = borde
    for indice, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(indice)].width = ancho
    if encabezados:
        ultima_columna = get_column_letter(len(encabezados))
        hoja.auto_filter.ref = f"A1:{ultima_columna}{max(1, len(datos) + 1)}"
        hoja.freeze_panes = "A2"
    for indice, formato in (formatos_numericos or {}).items():
        for fila in range(2, len(datos) + 2):
            hoja.cell(fila, indice + 1).number_format = formato
    hoja.sheet_properties.pageSetUpPr.fitToPage = True
    hoja.page_setup.fitToWidth = 1
    hoja.page_setup.fitToHeight = 0
    hoja.page_setup.orientation = "landscape"

    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


def generar_pdf_tabular(
    titulo: str,
    encabezados: Sequence[str],
    filas: Iterable[Sequence[Any]],
    pesos_columnas: Sequence[float],
    texto_resumen: str,
    columnas_horas: set[int] | None = None,
) -> bytes:
    """Crea un PDF apaisado y legible para un listado administrativo."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    datos = [list(fila) for fila in filas]
    columnas_horas = columnas_horas or set()
    estilos = getSampleStyleSheet()
    cuerpo_estilo = ParagraphStyle(
        "Tabla", parent=estilos["BodyText"], fontName="Helvetica", fontSize=7.5,
        leading=9, spaceAfter=0,
    )
    encabezado_estilo = ParagraphStyle(
        "Cabecera", parent=cuerpo_estilo, textColor=colors.white, fontName="Helvetica-Bold",
    )
    salida = BytesIO()
    documento = SimpleDocTemplate(
        salida, pagesize=landscape(A4),
        leftMargin=0.7 * cm, rightMargin=0.7 * cm,
        topMargin=0.8 * cm, bottomMargin=0.8 * cm,
    )
    elementos = [
        Paragraph(escape(titulo), estilos["Title"]),
        Paragraph(escape(texto_resumen), estilos["Normal"]),
        Spacer(1, 0.35 * cm),
    ]
    encabezado = [Paragraph(escape(titulo_columna), encabezado_estilo) for titulo_columna in encabezados]
    cuerpo = [
        [
            Paragraph(
                f"{float(valor):.1f} h" if indice in columnas_horas else escape(_texto(valor)),
                cuerpo_estilo,
            )
            for indice, valor in enumerate(fila)
        ]
        for fila in datos
    ] or [[Paragraph("No hay datos para exportar en el ámbito seleccionado.", cuerpo_estilo)] + [""] * (len(encabezados) - 1)]
    ancho_util = landscape(A4)[0] - 1.4 * cm
    total_pesos = sum(pesos_columnas) or len(encabezados)
    tabla = Table(
        [encabezado, *cuerpo], repeatRows=1,
        colWidths=[ancho_util * peso / total_pesos for peso in pesos_columnas],
    )
    tabla.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#B8C5D1")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F7FA")]),
    ]))
    elementos.append(tabla)
    documento.build(elementos)
    return salida.getvalue()


def generar_excel_adjudicaciones(filas: Iterable[dict[str, Any]]) -> bytes:
    datos = _filas_adjudicaciones(filas)
    return generar_excel_tabular(
        "Ordenación docente - adjudicaciones confirmadas", "Adjudicaciones",
        ENCABEZADOS_ADJUDICACIONES, datos,
        [24, 30, 38, 16, 15, 18, 17, 38, 29, 20],
        [("Líneas confirmadas", len(datos)), ("Horas confirmadas", sum(float(fila[-1]) for fila in datos))],
        {9: '0.0 "h"'},
    )


def generar_pdf_adjudicaciones(filas: Iterable[dict[str, Any]]) -> bytes:
    datos = _filas_adjudicaciones(filas)
    return generar_pdf_tabular(
        "Ordenación docente - adjudicaciones confirmadas", ENCABEZADOS_ADJUDICACIONES,
        datos, [2.3, 2.8, 3.8, 1.6, 1.5, 1.8, 1.7, 3.8, 2.8, 1.8],
        f"Generado el {datetime.now().strftime('%d/%m/%Y %H:%M')}. "
        f"{len(datos)} línea(s), {sum(float(fila[-1]) for fila in datos):.1f} h confirmadas.",
        {9},
    )
