"""Pruebas de los ficheros finales de adjudicaciones, sin tocar la base de datos."""
from io import BytesIO
import os
import sys
from zipfile import ZipFile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))


FILAS = [{
    "departamento": "Informática",
    "area": "Arquitectura y Tecnología de Computadores",
    "asignatura": "Estructura y Tecnología de Computadores",
    "codigo": "40151107",
    "grupo": "Grupo B",
    "tipo": "Teoría",
    "cuatrimestre": "1",
    "horario": "Lunes · 09:00-10:30",
    "profesor": "Pérez, Ana",
    "horas_confirmadas": 19.0,
}]


def test_excel_contiene_las_columnas_y_horas_confirmadas():
    from openpyxl import load_workbook
    from services.exportacion_adjudicaciones import generar_excel_adjudicaciones

    libro = load_workbook(BytesIO(generar_excel_adjudicaciones(FILAS)))
    assert libro.sheetnames == ["Resumen", "Adjudicaciones"]
    hoja = libro["Adjudicaciones"]
    assert hoja["C2"].value == "Estructura y Tecnología de Computadores"
    assert hoja["I2"].value == "Pérez, Ana"
    assert hoja["J2"].value == 19.0
    assert hoja.auto_filter.ref == "A1:J2"


def test_excel_no_incluye_tabla_xml_que_excel_deba_reparar():
    from services.exportacion_adjudicaciones import generar_excel_adjudicaciones

    with ZipFile(BytesIO(generar_excel_adjudicaciones(FILAS))) as archivo:
        assert not any(nombre.startswith("xl/tables/") for nombre in archivo.namelist())


def test_exportacion_tabular_admite_directorios_con_columnas_propias():
    from openpyxl import load_workbook
    from services.exportacion_adjudicaciones import generar_excel_tabular

    contenido = generar_excel_tabular(
        "Directorio de actividades", "Actividades", ["Asignatura", "Horas"],
        [["Bases de Datos", 6.0]], [32, 16], [("Actividades", 1)], {1: '0.0 "h"'},
    )
    libro = load_workbook(BytesIO(contenido))
    assert libro["Actividades"]["A2"].value == "Bases de Datos"
    assert libro["Actividades"]["B2"].value == 6.0


def test_pdf_se_genera_con_cabecera_valida():
    from services.exportacion_adjudicaciones import generar_pdf_adjudicaciones

    contenido = generar_pdf_adjudicaciones(FILAS)
    assert contenido.startswith(b"%PDF")
    assert len(contenido) > 1_000
